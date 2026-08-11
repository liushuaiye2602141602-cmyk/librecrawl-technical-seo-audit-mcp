"""Universal report — client safety, cross-client isolation, fixtures."""

import sys
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

docx = pytest.importorskip("docx")


def _build_docx(tmp_path, domain="https://site-a.example/", site_name="Site A",
                pages=228, items=None, audience="client"):
    from tests.phase12.test_universal_report_view import (
        _items, _metrics, _tasks,
    )
    from audit_rules.report_view import build_report_view
    from audit_rules.docx_report import build_universal_docx
    items = items if items is not None else _items()
    metrics = _metrics()
    view = build_report_view(
        items, _tasks(), metrics,
        manual_rows=[{"audit_id": 53, "rule": "Search intent",
                      "scope": "SITE", "status": "PENDING"}],
        technology_profile=None, technology_risks=[],
        domain=domain, audit_date="2026-08-11",
        pages_crawled=pages, audience=audience,
    )
    output = tmp_path / "universal.docx"
    build_universal_docx(
        str(output), view_model=view,
        items=items, task_rows=_tasks(), metrics=metrics,
        result_counts={"PASS": 36, "FAIL": 3, "WARNING": 4,
                       "OPPORTUNITY": 5, "MANUAL_REVIEW_REQUIRED": 8,
                       "UNKNOWN": 24},
        execution_counts={"EXECUTED_FULL": 38, "EXECUTED_PARTIAL": 11,
                          "NOT_CHECKED": 31, "NOT_APPLICABLE": 0},
        manual_rows=[{"audit_id": 53, "rule": "Search intent",
                      "scope": "SITE", "status": "PENDING"}],
        schema_distribution={},
        domain=domain, audit_date="2026-08-11",
        site_name=site_name, pages_crawled=pages,
    )
    return output


def _docx_text(path) -> str:
    from docx import Document
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells)


def test_shareable_safety_scan_passes_client_docx(tmp_path):
    from audit_rules.report_view import shareable_safety_scan
    path = _build_docx(tmp_path)
    violations = shareable_safety_scan(_docx_text(path))
    assert violations == [], f"client docx violations: {violations}"


def test_shareable_safety_scan_detects_bad_content():
    from audit_rules.report_view import shareable_safety_scan
    dirty = (
        "path C:\\Users\\me\\secret.py D:\\data\\x  token ghp_abcdefghijklmnopqrstuvwxyz0123456789 "
        "{'raw': 'dict'} Traceback TODO TBD example.com"
    )
    assert shareable_safety_scan(dirty)


def test_client_report_hides_internal_paths(tmp_path):
    path = _build_docx(tmp_path)
    text = _docx_text(path)
    assert "C:\\" not in text and "D:\\" not in text
    assert "\\Users\\" not in text


def test_client_report_hides_raw_dicts(tmp_path):
    path = _build_docx(tmp_path)
    text = _docx_text(path)
    assert "{'" not in text


def test_client_report_hides_tokens(tmp_path):
    path = _build_docx(tmp_path)
    text = _docx_text(path)
    for marker in ("ghp_", "AIza", "eyJ", "Bearer "):
        assert marker not in text


def test_client_report_hides_other_clients(tmp_path):
    path = _build_docx(tmp_path)
    text = _docx_text(path).lower()
    for marker in ("baolaipackaging", "gelgoogsort", "yashengcrafts"):
        assert marker not in text


def test_sequential_reports_do_not_leak_client_state(tmp_path):
    from audit_rules.report_view import build_report_view
    from tests.phase12.test_universal_report_view import (
        _items, _metrics, _tasks,
    )
    metrics_a = _metrics()
    metrics_a["score"] = 88.08
    items_a = _items()
    view_a = build_report_view(
        items_a, _tasks(), metrics_a, manual_rows=[], technology_profile=None,
        technology_risks=[], domain="https://site-a.example/",
        audit_date="2026-08-11", pages_crawled=228)

    items_b = _items()
    for item in items_b:
        item["check"] = item["check"].replace("Check", "CheckB")
        item["evidence"] = "evidence-b"
    metrics_b = _metrics()
    metrics_b["score"] = 55.0
    view_b = build_report_view(
        items_b, [], metrics_b, manual_rows=[], technology_profile=None,
        technology_risks=[], domain="https://site-b.example/",
        audit_date="2026-08-11", pages_crawled=33)

    text_b = "\n".join(view_b.management_summary) + " " + " ".join(
        f"{f['audit_id']}:{f['check']}" for f in view_b.key_findings)
    assert "site-a.example" not in text_b
    assert "88.08" not in text_b
    assert "228" not in text_b
    assert "55.0" in text_b
    assert "site-b.example" in view_b.domain or "site-b" in view_b.site_name


def test_current_run_metrics_only():
    from audit_rules.report_view import build_report_view
    from tests.phase12.test_universal_report_view import (
        _items, _metrics, _tasks,
    )
    view = build_report_view(
        _items(), _tasks(), _metrics(), manual_rows=[],
        technology_profile=None, technology_risks=[],
        domain="https://client-x.example/", audit_date="2026-08-11",
        pages_crawled=77)
    assert view.pages_crawled == 77
    assert view.report_metadata["pages_crawled"] == 77
