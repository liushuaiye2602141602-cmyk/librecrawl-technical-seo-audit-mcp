"""Universal Master SEO Diagnostic Report — DOCX structure tests."""

import sys
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

docx = pytest.importorskip("docx")


def _fixture(tmp_path, audience="client"):
    from tests.phase12.test_universal_report_view import (
        _items, _metrics, _tasks,
    )
    from audit_rules.report_view import build_report_view
    from audit_rules.docx_report import build_universal_docx
    items = _items()
    items[5]["representative"] = [
        "https://site-a.example/very-long-url-that-should-wrap-"
        + "x" * 120 + ".html"]
    task_rows = _tasks()
    metrics = _metrics()
    view = build_report_view(
        items, task_rows, metrics,
        manual_rows=[{"audit_id": 53, "rule": "Search intent"}],
        technology_profile={
            "detections": [{
                "category": "CMS", "technology_name": "WordPress",
                "status": "DETECTED", "confidence": "High",
                "version": "Unknown", "detection_sources": [],
                "limitations": [],
            }],
        },
        technology_risks=[],
        domain="https://site-a.example/", audit_date="2026-08-11",
        pages_crawled=228, audience=audience,
    )
    output = tmp_path / "universal.docx"
    build_universal_docx(
        str(output),
        view_model=view,
        items=items,
        task_rows=task_rows,
        metrics=metrics,
        result_counts={"PASS": 36, "FAIL": 3, "WARNING": 4,
                       "OPPORTUNITY": 5, "MANUAL_REVIEW_REQUIRED": 8,
                       "UNKNOWN": 24},
        execution_counts={"EXECUTED_FULL": 38, "EXECUTED_PARTIAL": 11,
                          "NOT_CHECKED": 31, "NOT_APPLICABLE": 0},
        manual_rows=[{"audit_id": 53}],
        schema_distribution={},
        technology_profile={
            "detections": [{
                "category": "CMS", "technology_name": "WordPress",
                "status": "DETECTED", "confidence": "High",
                "version": "Unknown", "detection_sources": [],
                "limitations": [],
            }],
        },
        technology_risks=[],
        domain="https://site-a.example/", audit_date="2026-08-11",
        site_name="Site A", pages_crawled=228,
    )
    return output


def _h1(doc):
    return [p.text for p in doc.paragraphs
            if p.style.name == "Heading 1"]


def test_universal_report_contains_management_summary(tmp_path):
    from docx import Document
    doc = Document(str(_fixture(tmp_path)))
    assert "Management Summary" in _h1(doc)


def test_universal_report_contains_key_findings(tmp_path):
    from docx import Document
    doc = Document(str(_fixture(tmp_path)))
    assert any("Key Findings" in h or "重点发现" in h for h in _h1(doc))


def test_universal_report_contains_remediation_plan(tmp_path):
    from docx import Document
    doc = Document(str(_fixture(tmp_path)))
    assert "Remediation Priority Plan" in _h1(doc)


def test_universal_report_contains_remediation_checklist(tmp_path):
    from docx import Document
    doc = Document(str(_fixture(tmp_path)))
    assert "Remediation Checklist" in _h1(doc)


def test_universal_report_contains_manual_review(tmp_path):
    from docx import Document
    doc = Document(str(_fixture(tmp_path)))
    assert any("Manual Review" in h for h in _h1(doc))


def test_universal_report_contains_recheck_workflow(tmp_path):
    from docx import Document
    doc = Document(str(_fixture(tmp_path)))
    assert "Acceptance & Recheck" in _h1(doc)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Verified" in text


def test_universal_report_contains_technology_profile(tmp_path):
    from docx import Document
    doc = Document(str(_fixture(tmp_path)))
    assert any("Technology Profile" in h for h in _h1(doc))


def test_docx_80_heading_regression(tmp_path):
    from docx import Document
    doc = Document(str(_fixture(tmp_path)))
    audits = [p for p in doc.paragraphs
              if p.style.name == "Heading 2" and p.text.startswith("AUDIT #")]
    assert len(audits) == 80
    assert audits[0].text == "AUDIT #01"
    assert audits[-1].text == "AUDIT #80"


def test_docx_long_urls_wrap(tmp_path):
    from docx import Document
    doc = Document(str(_fixture(tmp_path)))
    text = "\n".join(p.text for p in doc.paragraphs)
    long_url = ("https://site-a.example/very-long-url-that-should-wrap-"
                + "x" * 120 + ".html")
    assert long_url in text


def test_universal_report_cover_is_generic(tmp_path):
    from docx import Document
    doc = Document(str(_fixture(tmp_path)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Site A" in text
    assert "baolaipackaging" not in text.lower()
    assert "gelgoogsort" not in text.lower()
