"""DOCX client-report regression tests (FINAL WORD DELIVERY PATCH)."""

import sys
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

docx = pytest.importorskip("docx")


def _items(count: int = 80) -> list[dict]:
    items = []
    for audit_id in range(1, count + 1):
        items.append({
            "audit_id": audit_id,
            "category": "抓取与索引",
            "check": f"Check {audit_id}",
            "what_checked": f"What was checked for {audit_id}",
            "execution": "EXECUTED_FULL",
            "result": "PASS",
            "priority": "Medium",
            "confidence": "1.00",
            "data_source": "LibreCrawl",
            "actual_state": "Actual state text",
            "diagnosis": "Normal",
            "evidence": (
                "https://example.com/page-1.html — detected value | evidence text\n"
                "… 共 10 条，完整见 Detailed URL Findings"
            ),
            "affected_urls": 0,
            "representative": ["https://example.com/page-1.html",
                               "https://example.com/very-long-url-that-should-wrap-" + "x" * 120 + ".html"],
            "seo_impact": "No impact",
            "fix": "No remediation required.",
            "optional_maintenance": "Maintain current implementation. Literal example: <a href='URL'>",
            "owner": "SEO",
            "acceptance": "Acceptance criteria text",
            "limitations": "",
            "observed": "",
            "manual": {},
            "is_pass": True,
            "action_required": "None",
        })
    # #38 NOT_APPLICABLE with observation; #70 PARTIAL/UNKNOWN.
    items[37]["execution"] = "NOT_APPLICABLE"
    items[37]["result"] = "UNKNOWN"
    items[37]["observed"] = "Observation only; this WordPress-specific rule is not applicable."
    items[69]["execution"] = "EXECUTED_PARTIAL"
    items[69]["result"] = "UNKNOWN"
    items[69]["representative"] = ["SITE"]
    items[69]["evidence"] = (
        "SITE — Phase 2 form accessibility check is LIMITED — no DOM/form "
        "HTML available in crawl export. 5 likely form page(s) identified "
        "by URL patterns. | likely_form_urls=[...]")
    items[0]["acceptance"] = (
        "/robots.txt 返回 200；重要页面未被 Disallow；包含有效 Sitemap 声明。")
    items[1]["acceptance"] = (
        "Sitemap 中重要 URL 均为 200 + Indexable + Canonical；GSC/Bing 已成功读取。")
    return items


def _metrics() -> dict:
    return {
        "audit_rows": 80, "matrix_rows": 80, "coverage_rows": 80,
        "finding_rows": 909, "task_rows": 905, "manual_rows": 9,
        "performance_rows": 5, "pdf_pages": 0,
        "confirmed_remediation": 321, "optimization": 460,
        "data_required": 17, "manual_review_actions": 9,
        "p0": 0, "p1": 2, "p2": 319, "p3": 0,
        "score": 90.35, "coverage_pct": 64.29, "confidence_pct": 87.32,
    }


def _build(tmp_path) -> Path:
    from audit_rules.docx_report import build_docx
    output = tmp_path / "report.docx"
    build_docx(
        str(output),
        items=_items(),
        task_rows=[],
        metrics=_metrics(),
        result_counts={"PASS": 32, "FAIL": 2, "WARNING": 3, "OPPORTUNITY": 7,
                       "MANUAL_REVIEW_REQUIRED": 8, "UNKNOWN": 28},
        execution_counts={"EXECUTED_FULL": 34, "EXECUTED_PARTIAL": 11,
                          "NOT_CHECKED": 25, "NOT_APPLICABLE": 10},
        manual_rows=[{"audit_id": i} for i in (53, 54, 55, 56, 57, 71, 72, 73, 70)],
        schema_distribution={"Organization": 315, "BreadcrumbList": 315},
        domain="https://example.com/",
        site_name="Example Co",
        pages_crawled=315,
        http_ok_pages=315,
        audit_date="2026-08-09",
    )
    return output


def test_docx_contains_80_audits(tmp_path):
    from docx import Document
    doc = Document(str(_build(tmp_path)))
    audits = [
        p for p in doc.paragraphs
        if p.style.name == "Heading 2" and p.text.startswith("AUDIT #")
    ]
    assert len(audits) == 80
    assert audits[0].text == "AUDIT #01"
    assert audits[-1].text == "AUDIT #80"


def test_docx_summary_has_80_rows(tmp_path):
    from docx import Document
    doc = Document(str(_build(tmp_path)))
    table = max(doc.tables, key=lambda t: len(t.rows))
    assert len(table.rows) == 81
    assert len(table.columns) == 9
    assert table.rows[0].cells[0].text == "ID"


def test_docx_headings_are_real_styles(tmp_path):
    from docx import Document
    doc = Document(str(_build(tmp_path)))
    names = {p.style.name for p in doc.paragraphs}
    assert "Heading 1" in names and "Heading 2" in names
    style_names = {s.name for s in doc.styles}
    assert {"Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"} <= style_names


def test_docx_long_urls_wrap(tmp_path):
    from docx import Document
    doc = Document(str(_build(tmp_path)))
    text = "\n".join(p.text for p in doc.paragraphs)
    long_url = "https://example.com/very-long-url-that-should-wrap-" + "x" * 120 + ".html"
    assert long_url in text


def test_docx_literal_html_safe(tmp_path):
    from docx import Document
    doc = Document(str(_build(tmp_path)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "<a href='URL'>" in text
    rels = [v for v in doc.part.rels.values()
            if v.reltype.endswith("/hyperlink")]
    for rel in rels:
        if getattr(rel, "is_external", False):
            assert str(rel.target_ref).startswith(("http://", "https://"))


def test_docx_landscape_summary_section(tmp_path):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    doc = Document(str(_build(tmp_path)))
    orientations = [s.orientation for s in doc.sections]
    assert WD_ORIENT.LANDSCAPE in orientations


def test_docx_returns_to_portrait(tmp_path):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    doc = Document(str(_build(tmp_path)))
    orientations = [s.orientation for s in doc.sections]
    assert orientations[-1] == WD_ORIENT.PORTRAIT


def test_docx_manual_review_counts_match(tmp_path):
    from docx import Document
    doc = Document(str(_build(tmp_path)))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "9 条动作" in text or "9 条" in text


def test_docx_no_internal_engine_metadata(tmp_path):
    from docx import Document
    doc = Document(str(_build(tmp_path)))
    text = "\n".join(p.text for p in doc.paragraphs)
    for marker in ("pytest", "commit ", "git head", "CI green", "branch name"):
        assert marker not in text


def test_docx_final_metrics_consistent(tmp_path):
    from docx import Document
    doc = Document(str(_build(tmp_path)))
    table_text = "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    text = "\n".join(p.text for p in doc.paragraphs) + table_text
    assert "90.35" in text
    assert "64.29%" in text
    assert "87.32%" in text


def test_docx_roadmap_no_redundant_blank_break(tmp_path):
    """Section headings use page-break-before instead of a blank break
    paragraph, so LibreOffice/Word never render a header/footer-only page."""
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(str(_build(tmp_path)))
    paragraphs = doc.paragraphs
    roadmap = next(
        p for p in paragraphs
        if p.style.name == "Heading 1"
        and p.text == "30-Day Remediation Roadmap"
    )
    p_pr = roadmap._p.get_or_add_pPr()
    assert p_pr.find(qn("w:pageBreakBefore")) is not None
    # No empty paragraph immediately before the roadmap heading.
    index = paragraphs.index(roadmap)
    previous = paragraphs[index - 1]
    assert previous.text.strip() != ""


def _all_text(doc) -> str:
    table_text = "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    return "\n".join(p.text for p in doc.paragraphs) + "\n" + table_text


def test_docx_audit70_no_site_placeholder(tmp_path):
    from docx import Document
    doc = Document(str(_build(tmp_path)))
    text = _all_text(doc)
    assert "\nSITE\n" not in text
    assert "Rendered form DOM was not available." in text
    assert "Manual Validation Scope" in text


def test_docx_audit1_acceptance_updated(tmp_path):
    from docx import Document
    doc = Document(str(_build(tmp_path)))
    text = _all_text(doc)
    assert "如包含 Sitemap 声明，则地址有效" in text
    assert "包含有效 Sitemap 声明" not in text


def test_docx_audit2_partial_acceptance(tmp_path):
    from docx import Document
    doc = Document(str(_build(tmp_path)))
    text = _all_text(doc)
    assert "Automated/Crawl Acceptance" in text
    assert "External Validation" in text
    assert "GSC/Bing 已成功读取" not in text


def test_docx_cross_site_no_baolai_hardcode(tmp_path):
    """The DOCX builder must not leak first-site (Baolai) branding, page
    counts, language prefixes, or hashtag assumptions into a second site."""
    from docx import Document
    doc = Document(str(_build(tmp_path)))
    text = _all_text(doc)
    for marker in ("Baolai", "packaging", "/hashtag/", "零内部链接 15 页",
                   "309 页 / 3,011"):
        assert marker not in text, f"cross-site leak: {marker}"
    assert "Example Co" in text


def test_docx_header_uses_site_name(tmp_path):
    from docx import Document
    doc = Document(str(_build(tmp_path)))
    header_text = doc.sections[0].header.paragraphs[0].text
    assert "Example Co — 80-Item Master SEO Diagnostic Report" in header_text
    assert "Baolai" not in header_text
