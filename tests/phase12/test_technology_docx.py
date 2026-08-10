"""Task J 鈥?technology sections in the native DOCX report."""

import sys
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

docx = pytest.importorskip("docx")


def _profile(evidence_value=None, signal_type="meta_generator"):
    from audit_rules.technology.models import (
        TechnologyDetection, TechnologyEvidence, build_profile,
    )
    server_evidence = TechnologyEvidence(
        signal_type="response_header", signal_value="server=nginx",
        source_url="https://example.com/", source_scope="single_page",
        strength="strong", provenance="local_crawl")
    cms_evidence = TechnologyEvidence(
        signal_type=signal_type,
        signal_value=evidence_value or "WordPress",
        source_url="https://example.com/", source_scope="single_page",
        strength="strong", provenance="local_crawl")
    detections = [
        TechnologyDetection(
            category="Server", technology_name="nginx",
            technology_type="server", detection_sources=[server_evidence],
            confidence_score=0.9, confidence="High", status="DETECTED"),
        TechnologyDetection(
            category="CMS", technology_name="WordPress",
            technology_type="CMS", detection_sources=[cms_evidence],
            confidence_score=0.9, confidence="High", status="DETECTED"),
    ]
    return build_profile(
        detections, detector_version="1.0.0",
        signature_registry_version="1.0.0",
        source_url="https://example.com/", git_head="b" * 40)


def _risks():
    return [{
        "technology": "WordPress",
        "classification": "CONFIRMED_OBSERVATION",
        "mapped_audit_ids": [36, 37, 38, 39, 64, 65, 66, 67, 68, 69],
        "impact": "Provides evidence/applicability context for WordPress "
                  "audits; no direct score impact.",
        "recommended_action": "",
        "observation_status": "DETECTED",
        "observation_confidence": "High",
    }]


def _items(count: int = 80) -> list[dict]:
    items = []
    for audit_id in range(1, count + 1):
        items.append({
            "audit_id": audit_id,
            "category": "Crawling & Indexing",
            "check": f"Check {audit_id}",
            "what_checked": f"What was checked for {audit_id}",
            "execution": "EXECUTED_FULL",
            "result": "PASS",
            "priority": "Medium",
            "confidence": "1.00",
            "data_source": "LibreCrawl",
            "actual_state": "Actual state text",
            "diagnosis": "Normal",
            "evidence": "https://example.com/page-1.html 鈥?evidence",
            "affected_urls": 0,
            "representative": ["https://example.com/page-1.html"],
            "seo_impact": "No impact",
            "fix": "No remediation required.",
            "optional_maintenance": "",
            "owner": "SEO",
            "acceptance": "Acceptance criteria text",
            "limitations": "",
            "observed": "",
            "manual": {},
            "is_pass": True,
            "action_required": "None",
        })
    return items


def _metrics() -> dict:
    return {
        "audit_rows": 80, "matrix_rows": 80, "coverage_rows": 80,
        "finding_rows": 0, "task_rows": 0, "manual_rows": 0,
        "performance_rows": 0, "pdf_pages": 0,
        "confirmed_remediation": 0, "optimization": 0,
        "data_required": 0, "manual_review_actions": 0,
        "p0": 0, "p1": 0, "p2": 0, "p3": 0,
        "score": 90.35, "coverage_pct": 64.29, "confidence_pct": 87.32,
    }


def _build(tmp_path, profile=None, risks=None) -> Path:
    from audit_rules.docx_report import build_docx
    output = tmp_path / "report.docx"
    build_docx(
        str(output),
        items=_items(),
        task_rows=[],
        metrics=_metrics(),
        result_counts={"PASS": 80},
        execution_counts={"EXECUTED_FULL": 80},
        manual_rows=[],
        schema_distribution={},
        technology_profile=profile,
        technology_risks=risks,
    )
    return output


def _paragraph_text(doc) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


def _all_text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_technology_profile_present_in_docx(tmp_path):
    from docx import Document
    doc = Document(str(_build(tmp_path, profile=_profile())))
    text = _all_text(doc)
    assert "Website Technology Profile" in text
    assert "nginx" in text
    assert "WordPress" in text
    assert "High" in text


def test_technology_risks_present_in_docx(tmp_path):
    from docx import Document
    doc = Document(str(_build(tmp_path, profile=_profile(), risks=_risks())))
    text = _all_text(doc)
    assert "Technology Risks & Recommendations" in text
    assert "WordPress" in text
    assert "#36" in text


def test_technology_profile_contains_no_raw_internal_objects(tmp_path):
    from docx import Document
    doc = Document(str(_build(tmp_path, profile=_profile(), risks=_risks())))
    text = _all_text(doc)
    for marker in ("TechnologyEvidence", "object at 0x", "signal_type=",
                   "detection_sources", "TechnologyDetection"):
        assert marker not in text


def test_technology_docx_evidence_is_human_readable(tmp_path):
    from docx import Document
    doc = Document(str(_build(tmp_path, profile=_profile())))
    text = _all_text(doc)
    assert "Meta generator observed: WordPress" in text
    assert "Response header observed: server=nginx" in text


def test_technology_profile_has_status_column(tmp_path):
    from docx import Document
    doc = Document(str(_build(tmp_path, profile=_profile())))
    tech = next(t for t in doc.tables
                if [c.text for c in t.rows[0].cells][:2] ==
                ["Category", "Technology"])
    headers = [c.text for c in tech.rows[0].cells]
    assert headers == ["Category", "Technology", "Status", "Version",
                       "Confidence", "Evidence"]


def test_conflicting_status_wording_is_client_readable(tmp_path):
    from docx import Document
    profile = _profile()
    profile["detections"][1]["status"] = "CONFLICTING"
    doc = Document(str(_build(tmp_path, profile=profile)))
    text = _all_text(doc)
    assert (
        "Technology use is not confirmed; credible mutually-exclusive "
        "signals were observed." in text
    )


def test_technology_docx_contains_no_sensitive_header_values(tmp_path):
    from docx import Document
    profile = _profile(
        signal_type="response_header",
        evidence_value=(
            "server=nginx; Authorization: Bearer SECRET-TOKEN-12345"))
    doc = Document(str(_build(tmp_path, profile=profile)))
    text = _all_text(doc)
    assert "SECRET-TOKEN-12345" not in text
    assert "Bearer" not in text


def test_no_risk_sentence_when_no_confirmed_risk(tmp_path):
    from docx import Document
    doc = Document(str(_build(tmp_path, profile=_profile(), risks=[])))
    text = _paragraph_text(doc)
    assert (
        "No confirmed technology-specific risks were detected from the "
        "available observable evidence." in text
    )


def test_docx_still_contains_80_audits_after_technology_sections(tmp_path):
    from docx import Document
    doc = Document(str(_build(tmp_path, profile=_profile(), risks=_risks())))
    audits = [
        p for p in doc.paragraphs
        if p.style.name == "Heading 2" and p.text.startswith("AUDIT #")
    ]
    assert len(audits) == 80
    assert audits[0].text == "AUDIT #01"
    assert audits[-1].text == "AUDIT #80"
