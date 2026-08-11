"""FINAL UNIVERSAL CLIENT-WORDING DEDUP regressions."""

import sys
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

docx = pytest.importorskip("docx")


def _demo_view():
    from audit_rules.report_view import build_report_view
    from scripts.generate_universal_demo_report import (
        _build_items, _metrics, _tasks,
    )
    items, tasks = _build_items(), _tasks()
    view = build_report_view(
        items, tasks, _metrics(), manual_rows=[{"audit_id": 53}],
        technology_profile=None, technology_risks=[],
        domain="https://demo-client.example/", audit_date="2026-08-11",
        pages_crawled=228)
    return view, items, tasks


def _build_docx(tmp_path):
    from audit_rules.docx_report import build_universal_docx
    from scripts.generate_universal_demo_report import _metrics
    view, items, tasks = _demo_view()
    out = tmp_path / "u.docx"
    build_universal_docx(
        str(out), view_model=view, items=items, task_rows=tasks,
        metrics=_metrics(), result_counts=view.result_distribution,
        execution_counts=view.execution_distribution,
        manual_rows=[{"audit_id": 53}], schema_distribution={},
        domain="https://demo-client.example/", audit_date="2026-08-11",
        site_name="Demo Client", pages_crawled=228)
    return out


def _docx_text(path):
    from docx import Document
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells)


def test_potential_remediation_not_duplicated_in_primary_action():
    view, _, _ = _demo_view()
    for audit_id in (20, 34, 46, 53):
        item = next(i for i in view.audit_items if i["audit_id"] == audit_id)
        assert "Potential Remediation If Confirmed" not in item["what_to_do"]
        assert "Potential Remediation If Confirmed" not in item[
            "primary_action"]


def test_potential_remediation_label_rendered_once(tmp_path):
    path = _build_docx(tmp_path)
    text = _docx_text(path)
    # Exactly one label per applicable audit (4 audits), never doubled.
    label = "Potential Remediation If Confirmed:"
    assert text.count(label) == 4


def test_client_docx_contains_zero_double_potential_labels(tmp_path):
    path = _build_docx(tmp_path)
    text = _docx_text(path)
    assert (
        "Potential Remediation If Confirmed:\n"
        "Potential Remediation If Confirmed:" not in text)


def test_not_checked_why_it_matters_does_not_assume_failure():
    view, _, _ = _demo_view()
    for audit_id in (20, 34, 46):
        item = next(i for i in view.audit_items if i["audit_id"] == audit_id)
        assert "尚未验证" in item["why_it_matters"]
        assert "不应先修改网站" in item["why_it_matters"]
        assert "按整改建议处理后复测验收" not in item["why_it_matters"]


def test_manual_review_why_it_matters_does_not_assume_failure():
    view, _, _ = _demo_view()
    item = next(i for i in view.audit_items if i["audit_id"] == 53)
    assert "不能预设页面存在问题" in item["why_it_matters"]
    assert "先完成指定评审" in item["why_it_matters"]


def test_rule02_partial_pass_wording():
    view, _, _ = _demo_view()
    item = next(i for i in view.audit_items if i["audit_id"] == 2)
    assert item["execution"] == "EXECUTED_PARTIAL"
    assert item["result"] == "PASS"
    assert "GSC/Bing 外部提交/处理状态尚未验证" in item["why_it_matters"]
    assert "external validation remains outstanding" in item["why_it_matters"]
