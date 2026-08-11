"""Universal demo uses the REAL 80-rule definitions + client README + TOC."""

import sys
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

docx = pytest.importorskip("docx")


def _run_demo(monkeypatch):
    monkeypatch.setenv("DEMO_DATE", "2026-08-11-real")
    import scripts.generate_universal_demo_report as demo
    demo.main()
    return PROJECT_ROOT / "reports" / "universal-demo-2026-08-11-real"


def test_universal_demo_uses_real_rule_names(monkeypatch):
    out = _run_demo(monkeypatch)
    from audit_rules.registry import load_registry
    rules = {r.audit_id: r for r in load_registry()}
    import csv
    with open(out / "08_80_Rule_Coverage.csv", encoding="utf-8-sig",
              newline="") as f:
        rows = list(csv.DictReader(f))
    assert len(rows) == 80
    # The DOCX check names come from the registry titles.
    from docx import Document
    doc = Document(str(out / "01_Master_SEO_Diagnostic_Report.docx"))
    text = "\n".join(p.text for p in doc.paragraphs)
    for audit_id in (1, 6, 15, 29, 79):
        assert rules[audit_id].title in text


def test_universal_demo_uses_real_rule_categories(monkeypatch):
    out = _run_demo(monkeypatch)
    from audit_rules.registry import load_registry
    rules = {r.audit_id: r for r in load_registry()}
    from docx import Document
    doc = Document(str(out / "01_Master_SEO_Diagnostic_Report.docx"))
    text = "\n".join(p.text for p in doc.paragraphs)
    for audit_id in (1, 6, 15, 29, 79):
        assert rules[audit_id].category.value in text


def test_universal_demo_has_no_check_1_placeholder_names(monkeypatch):
    out = _run_demo(monkeypatch)
    from docx import Document
    doc = Document(str(out / "01_Master_SEO_Diagnostic_Report.docx"))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Check 1" not in text and "Check 80" not in text


def test_universal_demo_has_no_fixture_state_client_text(monkeypatch):
    out = _run_demo(monkeypatch)
    from docx import Document
    doc = Document(str(out / "01_Master_SEO_Diagnostic_Report.docx"))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Fixture state" not in text
    assert "recorded condition" not in text


def test_universal_demo_has_no_bare_evidence_placeholder(monkeypatch):
    out = _run_demo(monkeypatch)
    from docx import Document
    doc = Document(str(out / "01_Master_SEO_Diagnostic_Report.docx"))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "\nEvidence:\n证据" not in text
    assert "— evidence" not in text


def test_demo_docx_has_clickable_toc(monkeypatch):
    out = _run_demo(monkeypatch)
    from docx import Document
    doc = Document(str(out / "01_Master_SEO_Diagnostic_Report.docx"))
    paras = [p.text for p in doc.paragraphs]
    assert "Table of Contents" in paras
    # TOC anchors are bookmarks: verify hyperlink anchors exist in XML.
    xml = doc.element.xml
    for bookmark in ("ManagementSummary", "KeyFindings", "RemediationPlan",
                     "RemediationChecklist", "TechnicalAppendix"):
        assert bookmark in xml


def test_demo_readme_is_client_instructions(monkeypatch):
    out = _run_demo(monkeypatch)
    readme = (out / "README.txt").read_text(encoding="utf-8")
    for topic in ("开始阅读", "01_Master_SEO_Diagnostic_Report.docx",
                  "各文件用途", "PASS", "FAIL", "WARNING", "OPPORTUNITY",
                  "NOT_CHECKED", "NOT_APPLICABLE", "MANUAL_REVIEW_REQUIRED",
                  "整改工作流", "Checklist", "Acceptance & Recheck",
                  "不是 Google 官方评分"):
        assert topic in readme
    assert "C:\\" not in readme and "D:\\" not in readme


def test_demo_management_distribution_matches_80_rows(monkeypatch):
    out = _run_demo(monkeypatch)
    from audit_rules.report_consistency import validate_report_consistency
    from audit_rules.report_view import build_report_view
    from scripts.generate_universal_demo_report import (
        _build_items, _metrics, _tasks,
    )
    items, tasks = _build_items(), _tasks()
    view = build_report_view(
        items, tasks, _metrics(), manual_rows=[{"audit_id": 53}],
        technology_profile=None, technology_risks=[],
        domain="https://demo-client.example/", audit_date="2026-08-11",
        pages_crawled=228)
    violations = validate_report_consistency(view, items, tasks)
    assert violations == []
